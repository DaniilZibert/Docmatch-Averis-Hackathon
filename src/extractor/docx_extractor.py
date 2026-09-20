"""
OWNER: Person A.  Word SI / BL — 8 attachments.

The Word BL mirrors the real bilingual layout: a two-column table whose label cells
carry a Chinese gloss.

    ['Shipper (Principal or Seller) (发货人)', 'APRIL FINE PAPER TRADING | ON BEHALF OF ...']
    ['Gross Wt (kgs) (毛重 KGS)',              '243,588']
    ['PORT OF LOADING (装货港)',                'SINGAPORE']

Those labels never match an alias exactly — they resolve through the fuzzy pass in
field_aliases.match_field(), which is why rapidfuzz is a hard dependency.

Fields can also appear as plain paragraphs ("B/L NO.(提单号): SINF884969"), so both
`doc.tables` and `doc.paragraphs` are read.
"""

from __future__ import annotations

from pathlib import Path

from ..models import DocType, ExtractedDocument
from ._common import build_document, unreadable_document
from .txt_extractor import parse_label_lines


def extract_docx(path: Path, email_id: str, doc_type: DocType,
                 source_path: str) -> ExtractedDocument:
    import docx

    document = docx.Document(str(path))
    pairs: list[tuple[str, str]] = []
    text_lines: list[str] = []

    for table in document.tables:
        for row in table.rows:
            cells = [cell.text.strip() for cell in row.cells]
            if len(cells) < 2:
                continue
            # a cell holds "NAME\nADDRESS\nADDRESS"; the first line is the value,
            # the rest is the address (deliberately not compared).
            label = cells[0].splitlines()[0].strip() if cells[0] else ""
            value = cells[1].splitlines()[0].strip() if cells[1] else ""
            if label:
                pairs.append((label, value))
                text_lines.append(f"{label}: {value}")

    paragraph_text = "\n".join(p.text for p in document.paragraphs)
    pairs.extend(parse_label_lines(paragraph_text))
    text_lines.append(paragraph_text)

    if not pairs:
        return unreadable_document(email_id, doc_type, source_path,
                                   "document has no label/value rows")

    return build_document(pairs, email_id=email_id, doc_type=doc_type,
                          source_path=source_path, text="\n".join(text_lines))


__all__ = ["extract_docx"]
